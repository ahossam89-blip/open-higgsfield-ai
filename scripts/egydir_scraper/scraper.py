"""Scrape company listings from egydir.com (Egyptian business directory) into an .xlsx file.

Output columns: Company Name, Line of Business, Location, Has Website,
Website, Clients, Contact Number, Source URL.

Examples:
    # Crawl every "company" listing page and write companies.xlsx
    python -m scripts.egydir_scraper.scraper --output companies.xlsx

    # Also crawl factories/importers/exporters, cap at 10 pages each
    python -m scripts.egydir_scraper.scraper --types company factory importer exporter --max-pages 10

    # Best-effort client-name lookup on each company's own website (slow, unreliable)
    python -m scripts.egydir_scraper.scraper --fetch-clients

    # Sanity-check the parsing logic against the live site before a full run
    python -m scripts.egydir_scraper.scraper --dry-run

    # Rebuild the .xlsx from a checkpoint file without hitting the network again
    python -m scripts.egydir_scraper.scraper --export-only checkpoints/egydir.jsonl -o companies.xlsx

See README.md in this directory for setup notes, selector-tuning guidance and
the legal/rate-limiting caveats before running this against the real site.
"""

from __future__ import annotations

import argparse
import json
import logging
import random
import re
import time
import urllib.robotparser
from dataclasses import asdict, dataclass, field
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup

LOG = logging.getLogger("egydir_scraper")

USER_AGENT = (
    "EgydirCompanyResearchBot/1.0 (+contact: repository maintainer; "
    "respects robots.txt; polite crawl for a company data export)"
)

DEFAULT_BASE_URL = "https://www.egydir.com"
LISTING_TYPES = ("company", "factory", "importer", "exporter")

DETAIL_LINK_RE = re.compile(r"^/(?:en|ar)/(\d+)/")

# Egyptian phone numbers: landlines (0xx xxxx xxxx) and mobiles (01x xxxx xxxx),
# optionally with a +20 country code.
PHONE_RE = re.compile(r"(?:\+20|0020|0)\s?1?\d{1,2}[\s\-]?\d{3,4}[\s\-]?\d{3,4}\b")

SOCIAL_HOST_SUFFIXES = (
    "facebook.com", "fb.com", "twitter.com", "x.com", "instagram.com",
    "linkedin.com", "youtube.com", "youtu.be", "wa.me", "whatsapp.com",
    "google.com", "goo.gl", "maps.app.goo.gl", "t.me", "tiktok.com",
)

LOCATION_LABELS = ("address", "location", "city", "governorate", "area", "district")
PHONE_LABELS = ("phone", "tel", "telephone", "mobile", "fax", "contact")
WEBSITE_LABELS = ("website", "web site", "web", "url", "site")
SECTOR_LABELS = ("sector", "activity", "industry", "business", "category", "line of business")

CLIENTS_LINK_RE = re.compile(
    r"client|customer|partner|portfolio|our\s*work|"
    r"عملاء|شركاء",  # عملاء / شركاء
    re.IGNORECASE,
)
CLIENTS_CONTAINER_RE = re.compile(r"client|customer|partner", re.IGNORECASE)
GENERIC_ALT_WORDS = {"logo", "image", "img", "photo", "picture", "icon", "banner", "client", "customer"}


@dataclass
class CompanyRecord:
    id: str
    name: str
    line_of_business: str = ""
    location: str = ""
    has_website: str = "No"
    website: str = ""
    clients: str = ""
    contact_number: str = ""
    source_url: str = ""


@dataclass
class Fetcher:
    """Thin wrapper around requests with a shared session, delay, retries and
    a per-host robots.txt cache so the crawl stays polite by default."""

    delay: float = 1.5
    timeout: float = 15.0
    max_retries: int = 3
    session: requests.Session = field(default_factory=requests.Session)
    _robots_cache: dict = field(default_factory=dict)

    def __post_init__(self) -> None:
        self.session.headers.update({"User-Agent": USER_AGENT, "Accept-Language": "en,ar;q=0.8"})

    def _robots_allows(self, url: str) -> bool:
        parsed = urlparse(url)
        origin = f"{parsed.scheme}://{parsed.netloc}"
        rp = self._robots_cache.get(origin)
        if rp is None:
            rp = urllib.robotparser.RobotFileParser()
            rp.set_url(urljoin(origin, "/robots.txt"))
            try:
                rp.read()
            except Exception:
                rp = None  # unreadable robots.txt -> fail open, still rate-limited
            self._robots_cache[origin] = rp
        if rp is None:
            return True
        return rp.can_fetch(USER_AGENT, url)

    def get(self, url: str) -> requests.Response | None:
        if not self._robots_allows(url):
            LOG.warning("Skipping %s (disallowed by robots.txt)", url)
            return None
        last_exc: Exception | None = None
        for attempt in range(1, self.max_retries + 1):
            try:
                resp = self.session.get(url, timeout=self.timeout)
                if resp.status_code == 429 or resp.status_code >= 500:
                    raise requests.HTTPError(f"HTTP {resp.status_code}")
                resp.raise_for_status()
                time.sleep(self.delay + random.uniform(0, 0.5))
                return resp
            except Exception as exc:  # noqa: BLE001 - broad on purpose, network is unreliable
                last_exc = exc
                LOG.debug("GET %s failed (attempt %d/%d): %s", url, attempt, self.max_retries, exc)
                time.sleep(min(2 ** attempt, 10))
        LOG.warning("Giving up on %s: %s", url, last_exc)
        return None


def _clean(text: str | None) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def parse_listing_page(html: str, page_url: str) -> list[tuple[str, str, str]]:
    """Return [(id, name, detail_url), ...] found on a listing page."""
    soup = BeautifulSoup(html, "html.parser")
    seen: dict[str, tuple[str, str, str]] = {}
    for a in soup.find_all("a", href=True):
        m = DETAIL_LINK_RE.match(urlparse(a["href"]).path)
        if not m:
            continue
        company_id = m.group(1)
        name = _clean(a.get_text()) or _clean(a.get("title"))
        detail_url = urljoin(page_url, a["href"])
        if company_id in seen and (not name or len(name) < len(seen[company_id][1])):
            continue  # keep the more informative anchor text if a card has multiple links
        seen[company_id] = (company_id, name, detail_url)
    return list(seen.values())


def _extract_label_value_pairs(soup: BeautifulSoup) -> dict[str, str]:
    """Best-effort extraction of "label: value" pairs from common directory-page
    markup shapes (definition lists, table rows, icon+label+value rows)."""
    pairs: dict[str, str] = {}

    for dt in soup.find_all("dt"):
        dd = dt.find_next_sibling("dd")
        if dd:
            label = _clean(dt.get_text()).lower()
            value = _clean(dd.get_text())
            if label and value:
                pairs.setdefault(label, value)

    for row in soup.find_all("tr"):
        cells = row.find_all(["td", "th"])
        if len(cells) == 2:
            label = _clean(cells[0].get_text()).lower()
            value = _clean(cells[1].get_text())
            if label and value:
                pairs.setdefault(label, value)

    for li in soup.find_all(["li", "div", "p", "span"]):
        text = _clean(li.get_text())
        if ":" in text and 3 < len(text) < 200:
            label, _, value = text.partition(":")
            label, value = label.strip().lower(), value.strip()
            if label and value and len(label) < 30:
                pairs.setdefault(label, value)

    return pairs


def _match_label(pairs: dict[str, str], candidates: tuple[str, ...]) -> str:
    for label, value in pairs.items():
        if any(c in label for c in candidates):
            return value
    return ""


def _find_external_website(soup: BeautifulSoup, base_url: str) -> str:
    base_host = urlparse(base_url).netloc.lower()
    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        if not href.lower().startswith(("http://", "https://")):
            continue
        host = urlparse(href).netloc.lower().removeprefix("www.")
        if not host or host in base_host or base_host.endswith(host):
            continue
        if any(host == s or host.endswith("." + s) for s in SOCIAL_HOST_SUFFIXES):
            continue
        return href
    return ""


def parse_detail_page(html: str, url: str, base_url: str) -> dict:
    soup = BeautifulSoup(html, "html.parser")
    pairs = _extract_label_value_pairs(soup)
    full_text = _clean(soup.get_text(" "))

    location = _match_label(pairs, LOCATION_LABELS)
    line_of_business = _match_label(pairs, SECTOR_LABELS)

    phone = _match_label(pairs, PHONE_LABELS)
    tel_links = [
        a["href"].split(":", 1)[1]
        for a in soup.find_all("a", href=True)
        if a["href"].lower().startswith("tel:")
    ]
    phone_candidates = [_clean(p) for p in tel_links if _clean(p)]
    if not phone_candidates and phone:
        phone_candidates = [phone]
    if not phone_candidates:
        phone_candidates = sorted(set(PHONE_RE.findall(full_text)), key=full_text.find)
    phone_out = "; ".join(dict.fromkeys(p for p in phone_candidates if p))

    website = _match_label(pairs, WEBSITE_LABELS)
    if website and not website.lower().startswith(("http://", "https://")):
        # A labelled value that isn't a full URL (e.g. "www.example.com") is
        # still a real signal; a bare label with no domain-looking value is not.
        website = f"https://{website}" if "." in website else ""
    if not website:
        website = _find_external_website(soup, base_url)

    return {
        "line_of_business": line_of_business,
        "location": location,
        "contact_number": phone_out,
        "website": website,
    }


def find_client_names(fetcher: Fetcher, website: str, max_names: int = 30) -> list[str]:
    """Best-effort attempt to find a "clients" page on a company's own site and
    pull candidate client names from it. Heuristic and unreliable by nature —
    every company site is laid out differently. Returns [] on any failure."""
    try:
        resp = fetcher.get(website)
        if resp is None:
            return []
        soup = BeautifulSoup(resp.text, "html.parser")

        clients_url = website
        for a in soup.find_all("a", href=True):
            if CLIENTS_LINK_RE.search(_clean(a.get_text())):
                clients_url = urljoin(website, a["href"])
                break

        if clients_url != website:
            resp = fetcher.get(clients_url)
            if resp is None:
                return []
            soup = BeautifulSoup(resp.text, "html.parser")

        names: list[str] = []

        containers = soup.find_all(
            lambda tag: tag.has_attr("class") and any(CLIENTS_CONTAINER_RE.search(c) for c in tag["class"])
        ) + soup.find_all(
            lambda tag: tag.has_attr("id") and CLIENTS_CONTAINER_RE.search(tag["id"])
        )
        scope = containers or [soup]

        for container in scope:
            for img in container.find_all("img", alt=True):
                alt = _clean(img["alt"])
                if alt and alt.lower() not in GENERIC_ALT_WORDS and 1 < len(alt) <= 60:
                    names.append(alt)
            if not containers:
                continue
            for item in container.find_all(["li", "h3", "h4", "figcaption"]):
                text = _clean(item.get_text())
                if text and 1 < len(text) <= 60:
                    names.append(text)

        deduped = list(dict.fromkeys(names))
        return deduped[:max_names]
    except Exception as exc:  # noqa: BLE001
        LOG.debug("Client lookup failed for %s: %s", website, exc)
        return []


def iter_listing_pages(
    fetcher: Fetcher,
    base_url: str,
    listing_type: str,
    sector: int | None,
    city: int | None,
    start_page: int,
    max_pages: int | None,
):
    page = start_page
    pages_fetched = 0
    while max_pages is None or pages_fetched < max_pages:
        params = [f"page={page}"]
        if sector is not None:
            params.append(f"sector={sector}")
        if city is not None:
            params.append(f"city={city}")
        url = f"{base_url}/en/listings/{listing_type}?{'&'.join(params)}"
        resp = fetcher.get(url)
        pages_fetched += 1
        if resp is None:
            LOG.info("Stopping %s (sector=%s city=%s): request failed on page %d", listing_type, sector, city, page)
            return
        rows = parse_listing_page(resp.text, url)
        if not rows:
            LOG.info("Stopping %s (sector=%s city=%s): empty page %d", listing_type, sector, city, page)
            return
        yield url, rows
        page += 1


def load_checkpoint_ids(path: Path) -> set[str]:
    if not path.exists():
        return set()
    ids = set()
    with path.open("r", encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            try:
                ids.add(json.loads(line)["id"])
            except Exception:  # noqa: BLE001
                continue
    return ids


def append_checkpoint(path: Path, record: CompanyRecord) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as fh:
        fh.write(json.dumps(asdict(record), ensure_ascii=False) + "\n")


def write_excel(records: list[CompanyRecord], output_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Companies"

    headers = [
        "Company Name", "Line of Business", "Location", "Has Website",
        "Website", "Clients", "Contact Number", "Source URL",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for r in records:
        ws.append([
            r.name, r.line_of_business, r.location, r.has_website,
            r.website, r.clients, r.contact_number, r.source_url,
        ])

    # Fixed column widths tuned for the expected content (autofit isn't worth
    # the extra pass over every cell for a report this shape).
    widths = [28, 22, 20, 12, 30, 45, 22, 40]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def records_from_checkpoint(path: Path) -> list[CompanyRecord]:
    records = []
    with path.open("r", encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            records.append(CompanyRecord(**json.loads(line)))
    return records


def scrape(args: argparse.Namespace) -> list[CompanyRecord]:
    fetcher = Fetcher(delay=args.delay, timeout=args.timeout)
    checkpoint_path = Path(args.checkpoint)
    seen_ids = load_checkpoint_ids(checkpoint_path) if args.resume else set()
    if args.resume and seen_ids:
        LOG.info("Resuming: %d companies already in checkpoint", len(seen_ids))
    elif not args.resume and checkpoint_path.exists():
        checkpoint_path.unlink()

    sectors: list[int | None] = list(args.sectors) if args.sectors else [None]
    cities: list[int | None] = list(args.cities) if args.cities else [None]

    total_new = 0
    for listing_type in args.types:
        for sector in sectors:
            for city in cities:
                for page_url, rows in iter_listing_pages(
                    fetcher, args.base_url, listing_type, sector, city, args.start_page, args.max_pages
                ):
                    for company_id, name, detail_url in rows:
                        if company_id in seen_ids:
                            continue
                        if args.limit and total_new >= args.limit:
                            LOG.info("Reached --limit %d, stopping", args.limit)
                            return records_from_checkpoint(checkpoint_path)

                        detail_resp = fetcher.get(detail_url)
                        fields = (
                            parse_detail_page(detail_resp.text, detail_url, args.base_url)
                            if detail_resp is not None
                            else {}
                        )

                        record = CompanyRecord(
                            id=company_id,
                            name=name or f"Company #{company_id}",
                            line_of_business=fields.get("line_of_business", ""),
                            location=fields.get("location", ""),
                            website=fields.get("website", ""),
                            contact_number=fields.get("contact_number", ""),
                            source_url=detail_url,
                        )
                        record.has_website = "Yes" if record.website else "No"

                        if record.website and args.fetch_clients:
                            names = find_client_names(fetcher, record.website)
                            record.clients = "; ".join(names)
                        elif not record.website:
                            record.clients = ""

                        append_checkpoint(checkpoint_path, record)
                        seen_ids.add(company_id)
                        total_new += 1
                        LOG.info("[%d] %s (%s)", total_new, record.name, record.has_website)

    return records_from_checkpoint(checkpoint_path)


def dry_run(args: argparse.Namespace) -> None:
    """Fetch one listing page and one detail page, print what was parsed, and
    stop. Use this to sanity-check the selectors against the live site before
    committing to a full crawl."""
    fetcher = Fetcher(delay=args.delay, timeout=args.timeout)
    listing_type = args.types[0]
    url = f"{args.base_url}/en/listings/{listing_type}?page={args.start_page}"
    print(f"GET {url}")
    resp = fetcher.get(url)
    if resp is None:
        print("Request failed - check network access and --base-url.")
        return
    rows = parse_listing_page(resp.text, url)
    print(f"Found {len(rows)} company links on the listing page.")
    for company_id, name, detail_url in rows[:5]:
        print(f"  #{company_id}  {name!r}  ->  {detail_url}")
    if not rows:
        print("No company links matched. The site's markup may not match the "
              "DETAIL_LINK_RE pattern in scraper.py - inspect the HTML and adjust it.")
        return

    company_id, name, detail_url = rows[0]
    print(f"\nGET {detail_url}")
    detail_resp = fetcher.get(detail_url)
    if detail_resp is None:
        print("Detail page request failed.")
        return
    fields = parse_detail_page(detail_resp.text, detail_url, args.base_url)
    print(f"Parsed fields for {name!r}:")
    for key, value in fields.items():
        print(f"  {key}: {value!r}")


def build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--base-url", default=DEFAULT_BASE_URL, help="Site root (default: %(default)s)")
    p.add_argument("--types", nargs="+", choices=LISTING_TYPES, default=["company"],
                    help="Listing types to crawl (default: company)")
    p.add_argument("--sectors", nargs="+", type=int, default=None, help="Restrict to these sector IDs")
    p.add_argument("--cities", nargs="+", type=int, default=None, help="Restrict to these city IDs")
    p.add_argument("--start-page", type=int, default=1)
    p.add_argument("--max-pages", type=int, default=None, help="Max listing pages per type/sector/city combo")
    p.add_argument("--limit", type=int, default=None, help="Stop after this many new companies")
    p.add_argument("--delay", type=float, default=1.5, help="Seconds between requests (default: %(default)s)")
    p.add_argument("--timeout", type=float, default=15.0)
    p.add_argument("--fetch-clients", action="store_true",
                    help="Best-effort scrape of client names from each company's own website (slow, unreliable)")
    p.add_argument("--checkpoint", default="scripts/egydir_scraper/output/egydir_checkpoint.jsonl",
                    help="JSONL file written incrementally so a run can be resumed")
    p.add_argument("--resume", action="store_true", help="Skip companies already present in --checkpoint")
    p.add_argument("--output", "-o", default="scripts/egydir_scraper/output/egydir_companies.xlsx")
    p.add_argument("--export-only", metavar="CHECKPOINT_JSONL",
                    help="Skip scraping; just rebuild the .xlsx from an existing checkpoint file")
    p.add_argument("--dry-run", action="store_true",
                    help="Fetch one listing + one detail page, print parsed fields, then exit")
    p.add_argument("-v", "--verbose", action="store_true")
    return p


def main(argv: list[str] | None = None) -> None:
    args = build_arg_parser().parse_args(argv)
    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        datefmt="%H:%M:%S",
    )

    if args.export_only:
        records = records_from_checkpoint(Path(args.export_only))
        write_excel(records, Path(args.output))
        LOG.info("Wrote %d companies to %s", len(records), args.output)
        return

    if args.dry_run:
        dry_run(args)
        return

    try:
        records = scrape(args)
    except KeyboardInterrupt:
        LOG.warning("Interrupted - writing out what was collected so far")
        records = records_from_checkpoint(Path(args.checkpoint))

    write_excel(records, Path(args.output))
    LOG.info("Wrote %d companies to %s", len(records), args.output)


if __name__ == "__main__":
    main()
